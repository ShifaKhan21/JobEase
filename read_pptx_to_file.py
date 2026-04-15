from pptx import Presentation
import sys

def extract_text(pptx_path):
    print(f"Reading {pptx_path}")
    try:
        prs = Presentation(pptx_path)
    except Exception as e:
        print(f"Error reading file: {e}")
        return

    with open("extracted_pptx.txt", "w", encoding="utf-8") as f:
        for i, slide in enumerate(prs.slides):
            f.write(f"\n--- Slide {i+1} ---\n")
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    f.write(shape.text + "\n")

if __name__ == '__main__':
    extract_text(sys.argv[1])
