from pptx import Presentation
import sys

def extract_text(pptx_path):
    print(f"Reading {pptx_path}")
    try:
        prs = Presentation(pptx_path)
    except Exception as e:
        print(f"Error reading file: {e}")
        return

    for i, slide in enumerate(prs.slides):
        print(f"\n--- Slide {i+1} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                print(shape.text)

if __name__ == '__main__':
    extract_text(sys.argv[1])
